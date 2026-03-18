"""
Export-Modul: PDF-Berichte, Excel-Tabellen und Steuerberater-Paket
"""

import os
import shutil
import json
from datetime import datetime
from pathlib import Path
from typing import List, Dict, Optional

from .constants import STEUER_KATEGORIEN, DOK_TYPEN

# Optionale Imports
try:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (SimpleDocTemplate, Table, TableStyle,
                                    Paragraph, Spacer, HRFlowable,
                                    PageBreak, KeepTogether)
    from reportlab.lib.enums import TA_CENTER, TA_RIGHT, TA_LEFT
    HAS_REPORTLAB = True
except ImportError:
    HAS_REPORTLAB = False

try:
    import openpyxl
    from openpyxl.styles import (Font, Alignment, PatternFill, Border, Side,
                                  numbers)
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


class Exporter:

    def __init__(self):
        self.erstellt_am = datetime.now()

    # ──────────────────────────────────────────────────────────────
    # PDF-Bericht
    # ──────────────────────────────────────────────────────────────

    def export_pdf(self, dokumente: List[Dict], ziel_pfad: str,
                   steuerjahr: int, statistiken: Dict) -> bool:
        if not HAS_REPORTLAB:
            return self._export_pdf_txt_fallback(dokumente, ziel_pfad, steuerjahr, statistiken)

        try:
            doc = SimpleDocTemplate(
                ziel_pfad,
                pagesize=A4,
                rightMargin=2*cm, leftMargin=2*cm,
                topMargin=2.5*cm, bottomMargin=2*cm
            )
            styles = getSampleStyleSheet()

            # Eigene Stile
            titel_stil = ParagraphStyle(
                'Titel', parent=styles['Title'],
                fontSize=20, textColor=colors.HexColor('#1a3a5c'),
                spaceAfter=0.3*cm
            )
            h1_stil = ParagraphStyle(
                'H1', parent=styles['Heading1'],
                fontSize=14, textColor=colors.HexColor('#1a3a5c'),
                spaceBefore=0.5*cm, spaceAfter=0.2*cm
            )
            h2_stil = ParagraphStyle(
                'H2', parent=styles['Heading2'],
                fontSize=11, textColor=colors.HexColor('#2c5f8a'),
                spaceBefore=0.3*cm, spaceAfter=0.1*cm
            )
            normal_stil = ParagraphStyle(
                'Normal2', parent=styles['Normal'],
                fontSize=9, leading=12
            )
            fett_stil = ParagraphStyle(
                'Fett', parent=styles['Normal'],
                fontSize=10, fontName='Helvetica-Bold'
            )

            inhalt = []

            # ── Deckblatt ──
            inhalt.append(Spacer(1, 1*cm))
            inhalt.append(Paragraph(f"Steuererklärung {steuerjahr}", titel_stil))
            inhalt.append(Paragraph(
                f"Erstellt am {self.erstellt_am.strftime('%d.%m.%Y um %H:%M Uhr')}",
                normal_stil
            ))
            inhalt.append(HRFlowable(width="100%", thickness=2,
                                     color=colors.HexColor('#1a3a5c')))
            inhalt.append(Spacer(1, 0.5*cm))

            # ── Zusammenfassung ──
            inhalt.append(Paragraph("Übersicht", h1_stil))
            zusammen_daten = [
                ['Kennzahl', 'Wert'],
                ['Anzahl Dokumente', str(statistiken.get('total', 0))],
                ['Dokumente geprüft', str(sum(1 for d in dokumente if d.get('geprueft')))],
                ['Gesamtbetrag', f"{statistiken.get('gesamtsumme', 0):.2f} €"],
                ['Steuerjahr', str(steuerjahr)],
            ]
            t = Table(zusammen_daten, colWidths=[9*cm, 8*cm])
            t.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1a3a5c')),
                ('TEXTCOLOR',  (0,0), (-1,0), colors.white),
                ('FONTNAME',   (0,0), (-1,0), 'Helvetica-Bold'),
                ('FONTSIZE',   (0,0), (-1,-1), 9),
                ('ROWBACKGROUNDS', (0,1), (-1,-1),
                 [colors.HexColor('#f0f4f8'), colors.white]),
                ('GRID',       (0,0), (-1,-1), 0.5, colors.HexColor('#ccc')),
                ('ALIGN',      (1,0), (1,-1), 'RIGHT'),
                ('TOPPADDING', (0,0), (-1,-1), 5),
                ('BOTTOMPADDING', (0,0), (-1,-1), 5),
            ]))
            inhalt.append(t)
            inhalt.append(Spacer(1, 0.5*cm))

            # ── Steuer-Kategorien-Übersicht ──
            inhalt.append(Paragraph("Steuerliche Kategorien", h1_stil))

            for haupt_key, haupt_data in STEUER_KATEGORIEN.items():
                haupt_stats = statistiken.get('nach_steuer', {}).get(haupt_key, {})
                if not haupt_stats:
                    continue

                gesamt_haupt = sum(v.get('summe', 0) for v in haupt_stats.values())
                inhalt.append(Paragraph(
                    f"{haupt_data['label']}  —  Gesamt: {gesamt_haupt:.2f} €",
                    h2_stil
                ))

                kat_daten = [['Unterkategorie', 'Anzahl', 'Summe (€)']]
                for sub_key, sub_data in haupt_data['subcats'].items():
                    sub_stats = haupt_stats.get(sub_key, {})
                    if sub_stats:
                        kat_daten.append([
                            sub_data['label'],
                            str(sub_stats.get('anzahl', 0)),
                            f"{sub_stats.get('summe', 0):.2f}"
                        ])

                if len(kat_daten) > 1:
                    t2 = Table(kat_daten, colWidths=[11*cm, 2.5*cm, 3.5*cm])
                    t2.setStyle(TableStyle([
                        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#2c5f8a')),
                        ('TEXTCOLOR',  (0,0), (-1,0), colors.white),
                        ('FONTNAME',   (0,0), (-1,0), 'Helvetica-Bold'),
                        ('FONTSIZE',   (0,0), (-1,-1), 8),
                        ('ROWBACKGROUNDS', (0,1), (-1,-1),
                         [colors.HexColor('#e8f0f8'), colors.white]),
                        ('GRID',       (0,0), (-1,-1), 0.3, colors.HexColor('#bbb')),
                        ('ALIGN',      (1,0), (2,-1), 'RIGHT'),
                        ('TOPPADDING', (0,0), (-1,-1), 4),
                        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
                    ]))
                    inhalt.append(t2)
                    inhalt.append(Spacer(1, 0.2*cm))

            # ── Dokument-Liste ──
            inhalt.append(PageBreak())
            inhalt.append(Paragraph("Alle Dokumente", h1_stil))

            # Gruppiert nach Typ
            nach_typ: Dict[str, List] = {}
            for d in dokumente:
                typ = d.get('dok_typ', 'sonstiges')
                nach_typ.setdefault(typ, []).append(d)

            for typ, docs in sorted(nach_typ.items()):
                typ_label = DOK_TYPEN.get(typ, typ.capitalize())
                inhalt.append(Paragraph(f"{typ_label} ({len(docs)})", h2_stil))

                dok_daten = [['Dateiname', 'Datum', 'Betrag (€)', 'Lieferant', '✓']]
                for d in docs:
                    dok_daten.append([
                        Paragraph(d.get('dateiname', '')[:50], normal_stil),
                        d.get('datum', '')[:10] if d.get('datum') else '—',
                        f"{d.get('betrag', 0):.2f}" if d.get('betrag') else '—',
                        (d.get('lieferant') or '—')[:30],
                        '✓' if d.get('geprueft') else ''
                    ])

                if len(dok_daten) > 1:
                    t3 = Table(dok_daten,
                               colWidths=[6.5*cm, 2.5*cm, 2.5*cm, 4*cm, 1*cm])
                    t3.setStyle(TableStyle([
                        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#374151')),
                        ('TEXTCOLOR',  (0,0), (-1,0), colors.white),
                        ('FONTNAME',   (0,0), (-1,0), 'Helvetica-Bold'),
                        ('FONTSIZE',   (0,0), (-1,-1), 7),
                        ('ROWBACKGROUNDS', (0,1), (-1,-1),
                         [colors.HexColor('#f9fafb'), colors.white]),
                        ('GRID',       (0,0), (-1,-1), 0.2, colors.HexColor('#e5e7eb')),
                        ('ALIGN',      (2,0), (2,-1), 'RIGHT'),
                        ('ALIGN',      (4,0), (4,-1), 'CENTER'),
                        ('TOPPADDING', (0,0), (-1,-1), 3),
                        ('BOTTOMPADDING', (0,0), (-1,-1), 3),
                        ('WORDWRAP',   (0,1), (0,-1), True),
                    ]))
                    inhalt.append(KeepTogether([t3, Spacer(1, 0.3*cm)]))

            # ── EÜR ──
            inhalt.append(PageBreak())
            inhalt.append(Paragraph(
                f"Einnahmen-Überschuss-Rechnung (EÜR) {steuerjahr}", h1_stil))
            inhalt.append(self._erstelle_eur_tabelle(dokumente, styles, normal_stil))

            doc.build(inhalt, onFirstPage=self._seiten_rahmen,
                      onLaterPages=self._seiten_rahmen)
            return True

        except Exception as e:
            print(f"PDF-Fehler: {e}")
            return self._export_pdf_txt_fallback(dokumente, ziel_pfad, steuerjahr, statistiken)

    def _erstelle_eur_tabelle(self, dokumente, styles, normal_stil):
        if not HAS_REPORTLAB:
            return Spacer(1, 0)

        einnahmen = sum(
            d.get('betrag', 0) or 0
            for d in dokumente
            if d.get('steuer_haupt') == 'EINNAHMEN'
        )
        ausgaben = sum(
            d.get('betrag', 0) or 0
            for d in dokumente
            if d.get('steuer_haupt') in (
                'WERBUNGSKOSTEN', 'BETRIEBSAUSGABEN',
                'SONDERAUSGABEN', 'AUSSERGEWOEHNLICH', 'HAUSHALTSNAHE'
            )
        )
        gewinn = einnahmen - ausgaben

        eur_daten = [
            ['Position', 'Betrag (€)'],
            ['Einnahmen gesamt', f'{einnahmen:.2f}'],
            ['Ausgaben gesamt', f'{ausgaben:.2f}'],
            ['', ''],
            ['Überschuss / Gewinn', f'{gewinn:.2f}'],
        ]

        h = colors.HexColor
        t = Table(eur_daten, colWidths=[12*cm, 5*cm])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), h('#1a3a5c')),
            ('TEXTCOLOR',  (0,0), (-1,0), colors.white),
            ('FONTNAME',   (0,0), (-1,0), 'Helvetica-Bold'),
            ('FONTNAME',   (0,4), (-1,4), 'Helvetica-Bold'),
            ('BACKGROUND', (0,4), (-1,4), h('#e8f5e9')),
            ('FONTSIZE',   (0,0), (-1,-1), 10),
            ('GRID',       (0,0), (-1,-1), 0.5, h('#ccc')),
            ('ALIGN',      (1,0), (1,-1), 'RIGHT'),
            ('TOPPADDING', (0,0), (-1,-1), 6),
            ('BOTTOMPADDING', (0,0), (-1,-1), 6),
        ]))
        return t

    def _seiten_rahmen(self, canvas, doc):
        canvas.saveState()
        canvas.setFont('Helvetica', 8)
        canvas.setFillColor(colors.HexColor('#666'))
        w, h = A4
        canvas.drawString(2*cm, 1.5*cm,
                          f"SteuerApp – Steuererklärung | Erstellt: "
                          f"{self.erstellt_am.strftime('%d.%m.%Y')}")
        canvas.drawRightString(w - 2*cm, 1.5*cm,
                               f"Seite {doc.page}")
        canvas.restoreState()

    def _export_pdf_txt_fallback(self, dokumente, ziel_pfad, steuerjahr, statistiken):
        """Fallback: Einfache Textdatei wenn reportlab fehlt"""
        txt_pfad = ziel_pfad.replace('.pdf', '_bericht.txt')
        with open(txt_pfad, 'w', encoding='utf-8') as f:
            f.write(f"STEUERERKLÄRUNG {steuerjahr}\n")
            f.write("=" * 60 + "\n\n")
            f.write(f"Erstellt: {self.erstellt_am.strftime('%d.%m.%Y %H:%M')}\n")
            f.write(f"Anzahl Dokumente: {statistiken.get('total', 0)}\n")
            f.write(f"Gesamtbetrag: {statistiken.get('gesamtsumme', 0):.2f} EUR\n\n")

            f.write("DOKUMENTE\n" + "-" * 40 + "\n")
            for d in dokumente:
                f.write(f"  {d.get('dateiname', 'Unbekannt')}\n")
                f.write(f"    Typ: {DOK_TYPEN.get(d.get('dok_typ',''), d.get('dok_typ',''))}\n")
                if d.get('betrag'):
                    f.write(f"    Betrag: {d['betrag']:.2f} EUR\n")
                if d.get('datum'):
                    f.write(f"    Datum: {d['datum']}\n")
                if d.get('lieferant'):
                    f.write(f"    Von: {d['lieferant']}\n")
                f.write("\n")
        return True

    # ──────────────────────────────────────────────────────────────
    # Excel-Export
    # ──────────────────────────────────────────────────────────────

    def export_excel(self, dokumente: List[Dict], ziel_pfad: str,
                     steuerjahr: int, statistiken: Dict) -> bool:
        if not HAS_OPENPYXL:
            return self._export_csv_fallback(dokumente, ziel_pfad)

        try:
            wb = openpyxl.Workbook()
            self._erstelle_blatt_dokumente(wb, dokumente)
            self._erstelle_blatt_kategorien(wb, statistiken)
            self._erstelle_blatt_eur(wb, dokumente, steuerjahr)
            self._erstelle_blatt_steuerberater(wb, dokumente, steuerjahr, statistiken)

            # Standard-Blatt entfernen
            if 'Sheet' in wb.sheetnames:
                del wb['Sheet']

            wb.save(ziel_pfad)
            return True
        except Exception as e:
            print(f"Excel-Fehler: {e}")
            return self._export_csv_fallback(dokumente, ziel_pfad)

    def _erstelle_blatt_dokumente(self, wb, dokumente):
        ws = wb.create_sheet("Alle Dokumente")
        ws.sheet_view.showGridLines = True

        kopf = ['Nr.', 'Dateiname', 'Typ', 'Datum', 'Betrag (€)',
                'Lieferant/Von', 'Steuer-Kategorie', 'Unterkategorie',
                'Geprüft', 'Notiz', 'Pfad']

        h_fill = PatternFill('solid', fgColor='1a3a5c')
        h_font = Font(color='FFFFFF', bold=True, size=10)

        for col, h in enumerate(kopf, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.fill = h_fill
            cell.font = h_font
            cell.alignment = Alignment(horizontal='center', vertical='center')

        for i, d in enumerate(dokumente, 1):
            row = i + 1
            ws.cell(row=row, column=1, value=i)
            ws.cell(row=row, column=2, value=d.get('dateiname', ''))
            ws.cell(row=row, column=3,
                    value=DOK_TYPEN.get(d.get('dok_typ',''), d.get('dok_typ','')))
            ws.cell(row=row, column=4, value=d.get('datum', '')[:10] if d.get('datum') else '')
            betrag_cell = ws.cell(row=row, column=5, value=d.get('betrag') or 0)
            betrag_cell.number_format = '#,##0.00 €'
            ws.cell(row=row, column=6, value=d.get('lieferant', ''))
            ws.cell(row=row, column=7, value=d.get('steuer_haupt', ''))
            ws.cell(row=row, column=8, value=d.get('steuer_sub', ''))
            ws.cell(row=row, column=9, value='Ja' if d.get('geprueft') else 'Nein')
            ws.cell(row=row, column=10, value=d.get('notiz', ''))
            ws.cell(row=row, column=11, value=d.get('dateipfad', ''))

            # Abwechselnde Zeilenfarbe
            if i % 2 == 0:
                fill = PatternFill('solid', fgColor='e8f0f8')
                for col in range(1, 12):
                    ws.cell(row=row, column=col).fill = fill

        # Spaltenbreiten
        breiten = [5, 35, 20, 12, 12, 25, 20, 20, 8, 20, 50]
        for col, b in enumerate(breiten, 1):
            ws.column_dimensions[get_column_letter(col)].width = b

        ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes = 'A2'

    def _erstelle_blatt_kategorien(self, wb, statistiken):
        ws = wb.create_sheet("Steuer-Kategorien")

        ws['A1'] = 'Steuerliche Kategorien – Übersicht'
        ws['A1'].font = Font(bold=True, size=14, color='1a3a5c')
        ws.merge_cells('A1:D1')

        zeile = 3
        for haupt_key, haupt_data in STEUER_KATEGORIEN.items():
            haupt_stats = statistiken.get('nach_steuer', {}).get(haupt_key, {})
            gesamt = sum(v.get('summe', 0) for v in haupt_stats.values())

            ws.cell(row=zeile, column=1, value=haupt_data['label'])
            ws.cell(row=zeile, column=1).font = Font(bold=True, size=11, color='1a3a5c')
            ws.cell(row=zeile, column=3, value=gesamt)
            ws.cell(row=zeile, column=3).number_format = '#,##0.00 €'
            ws.cell(row=zeile, column=3).font = Font(bold=True)
            zeile += 1

            for sub_key, sub_data in haupt_data['subcats'].items():
                sub_stats = haupt_stats.get(sub_key, {})
                ws.cell(row=zeile, column=2, value=sub_data['label'])
                ws.cell(row=zeile, column=3, value=sub_stats.get('anzahl', 0))
                ws.cell(row=zeile, column=4, value=sub_stats.get('summe', 0))
                ws.cell(row=zeile, column=4).number_format = '#,##0.00 €'
                zeile += 1

            zeile += 1

        ws.column_dimensions['A'].width = 5
        ws.column_dimensions['B'].width = 40
        ws.column_dimensions['C'].width = 12
        ws.column_dimensions['D'].width = 15

    def _erstelle_blatt_eur(self, wb, dokumente, steuerjahr):
        ws = wb.create_sheet("EÜR")

        ws['A1'] = f'Einnahmen-Überschuss-Rechnung {steuerjahr}'
        ws['A1'].font = Font(bold=True, size=14, color='1a3a5c')
        ws.merge_cells('A1:C1')

        zeile = 3
        ws.cell(row=zeile, column=1, value='Position').font = Font(bold=True)
        ws.cell(row=zeile, column=2, value='Anzahl Belege').font = Font(bold=True)
        ws.cell(row=zeile, column=3, value='Betrag (€)').font = Font(bold=True)
        zeile += 1

        einnahmen = ausgaben = 0

        ws.cell(row=zeile, column=1, value='EINNAHMEN').font = Font(bold=True, color='1a5c1a')
        zeile += 1

        haupt_einnahmen = [
            ('EINNAHMEN', 'gehalt', 'Gehalt/Lohn'),
            ('EINNAHMEN', 'selbstaendig', 'Selbständige Tätigkeit'),
            ('EINNAHMEN', 'gewerbe', 'Gewerbebetrieb'),
            ('EINNAHMEN', 'vermietung', 'Vermietung & Verpachtung'),
            ('EINNAHMEN', 'kapital', 'Kapitalerträge'),
            ('EINNAHMEN', 'rente', 'Renten'),
            ('EINNAHMEN', 'sonstige_ein', 'Sonstige Einnahmen'),
        ]

        for haupt, sub, label in haupt_einnahmen:
            docs = [d for d in dokumente
                    if d.get('steuer_haupt') == haupt and d.get('steuer_sub') == sub]
            betrag = sum(d.get('betrag', 0) or 0 for d in docs)
            einnahmen += betrag
            ws.cell(row=zeile, column=1, value=f'  {label}')
            ws.cell(row=zeile, column=2, value=len(docs))
            ws.cell(row=zeile, column=3, value=betrag).number_format = '#,##0.00 €'
            zeile += 1

        ws.cell(row=zeile, column=1, value='Einnahmen gesamt').font = Font(bold=True)
        ws.cell(row=zeile, column=3, value=einnahmen)
        ws.cell(row=zeile, column=3).number_format = '#,##0.00 €'
        ws.cell(row=zeile, column=3).font = Font(bold=True)
        zeile += 2

        ws.cell(row=zeile, column=1, value='AUSGABEN / ABZÜGE').font = Font(bold=True, color='5c1a1a')
        zeile += 1

        ausgaben_kategorien = [
            ('WERBUNGSKOSTEN', None, 'Werbungskosten gesamt'),
            ('BETRIEBSAUSGABEN', None, 'Betriebsausgaben gesamt'),
            ('SONDERAUSGABEN', None, 'Sonderausgaben gesamt'),
            ('AUSSERGEWOEHNLICH', None, 'Außergewöhnl. Belastungen'),
            ('HAUSHALTSNAHE', None, 'Haushaltsnahe Ausgaben'),
        ]

        for haupt, _, label in ausgaben_kategorien:
            docs = [d for d in dokumente if d.get('steuer_haupt') == haupt]
            betrag = sum(d.get('betrag', 0) or 0 for d in docs)
            ausgaben += betrag
            ws.cell(row=zeile, column=1, value=f'  {label}')
            ws.cell(row=zeile, column=2, value=len(docs))
            ws.cell(row=zeile, column=3, value=betrag).number_format = '#,##0.00 €'
            zeile += 1

        ws.cell(row=zeile, column=1, value='Ausgaben/Abzüge gesamt').font = Font(bold=True)
        ws.cell(row=zeile, column=3, value=ausgaben)
        ws.cell(row=zeile, column=3).number_format = '#,##0.00 €'
        ws.cell(row=zeile, column=3).font = Font(bold=True)
        zeile += 2

        gewinn = einnahmen - ausgaben
        ws.cell(row=zeile, column=1, value='ÜBERSCHUSS / GEWINN').font = Font(bold=True, size=12)
        ws.cell(row=zeile, column=3, value=gewinn)
        ws.cell(row=zeile, column=3).number_format = '#,##0.00 €'
        ws.cell(row=zeile, column=3).font = Font(bold=True, size=12)
        if gewinn >= 0:
            ws.cell(row=zeile, column=3).fill = PatternFill('solid', fgColor='c8e6c9')
        else:
            ws.cell(row=zeile, column=3).fill = PatternFill('solid', fgColor='ffcdd2')

        ws.column_dimensions['A'].width = 35
        ws.column_dimensions['B'].width = 14
        ws.column_dimensions['C'].width = 16

    def _erstelle_blatt_steuerberater(self, wb, dokumente, steuerjahr, statistiken):
        ws = wb.create_sheet("Steuerberater-Übergabe")

        ws['A1'] = f'Übergabeprotokoll für Steuerberater – Steuerjahr {steuerjahr}'
        ws['A1'].font = Font(bold=True, size=13, color='1a3a5c')
        ws.merge_cells('A1:E1')

        ws['A3'] = f'Erstellt am: {datetime.now().strftime("%d.%m.%Y %H:%M")}'
        ws['A4'] = f'Anzahl Dokumente: {statistiken.get("total", 0)}'
        ws['A5'] = f'Dokumente geprüft: {sum(1 for d in dokumente if d.get("geprueft"))}'
        ws['A6'] = f'Gesamtbetrag: {statistiken.get("gesamtsumme", 0):.2f} EUR'

        zeile = 9
        ws.cell(row=zeile, column=1, value='CHECKLISTE').font = Font(bold=True, size=11)
        zeile += 1

        checkliste = [
            'Lohnsteuerbescheinigung(en)',
            'Anlage N (Werbungskosten)',
            'Sonderausgaben (Versicherungen)',
            'Außergewöhnliche Belastungen',
            'Haushaltsnahe Dienstleistungen',
            'Kapitalerträge (Anlage KAP)',
            'Vermietung/Verpachtung (Anlage V)',
            'Belege geordnet und nummeriert',
            'Steuernummer / Identifikationsnummer',
        ]

        for punkt in checkliste:
            ws.cell(row=zeile, column=1, value='☐')
            ws.cell(row=zeile, column=2, value=punkt)
            zeile += 1

        ws.column_dimensions['A'].width = 5
        ws.column_dimensions['B'].width = 45

    def _export_csv_fallback(self, dokumente, ziel_pfad):
        """CSV-Fallback wenn openpyxl fehlt"""
        csv_pfad = ziel_pfad.replace('.xlsx', '.csv')
        try:
            import csv
            with open(csv_pfad, 'w', newline='', encoding='utf-8-sig') as f:
                writer = csv.DictWriter(f, fieldnames=[
                    'dateiname', 'dok_typ', 'datum', 'betrag', 'lieferant',
                    'steuer_haupt', 'steuer_sub', 'geprueft', 'notiz'
                ], extrasaction='ignore')
                writer.writeheader()
                writer.writerows(dokumente)
            return True
        except Exception:
            return False

    # ──────────────────────────────────────────────────────────────
    # Steuerberater-Paket (Ordnerstruktur)
    # ──────────────────────────────────────────────────────────────

    def erstelle_steuerberater_paket(self, dokumente: List[Dict],
                                     ziel_ordner: str, steuerjahr: int,
                                     statistiken: Dict) -> str:
        """
        Erstellt einen sortierten Ordner mit allen Belegen + Berichte.
        """
        basis = Path(ziel_ordner) / f"Steuererklarung_{steuerjahr}"
        basis.mkdir(parents=True, exist_ok=True)

        # Unterordner erstellen
        unterordner = {
            'rechnung':           '02_Rechnungen',
            'quittung':           '03_Quittungen_Kassenbons',
            'gehaltsabrechnung':  '04_Gehaltsabrechnungen',
            'lohnsteuerbescheid': '05_Lohnsteuerbescheinigungen',
            'steuerbescheid':     '06_Steuerbescheide',
            'versicherung':       '07_Versicherungen',
            'kontoauszug':        '08_Kontoauszüge',
            'spende':             '09_Spendenquittungen',
            'arzt':               '10_Arzt_Medizin',
            'fahrtkosten':        '11_Fahrtkosten',
            'fortbildung':        '12_Fortbildung_Weiterbildung',
            'buero':              '13_Büro_Arbeitsmittel',
            'miete':              '14_Miete_Wohnen',
            'sonstiges':          '15_Sonstiges',
        }

        for ordner in unterordner.values():
            (basis / ordner).mkdir(exist_ok=True)

        (basis / '01_Berichte').mkdir(exist_ok=True)

        # Dateien kopieren
        kopiert = 0
        nicht_gefunden = []

        for d in dokumente:
            src = d.get('dateipfad', '')
            if not src or not os.path.exists(src):
                nicht_gefunden.append(d.get('dateiname', 'Unbekannt'))
                continue

            typ = d.get('dok_typ', 'sonstiges')
            ordner = unterordner.get(typ, '15_Sonstiges')

            # Dateiname mit Datum und Betrag anreichern
            datum = (d.get('datum', '')[:10] or 'kein-datum').replace('/', '-')
            betrag = f"_{d['betrag']:.2f}EUR" if d.get('betrag') else ''
            lieferant = (d.get('lieferant') or '').replace('/', '-')[:20]
            lieferant = ''.join(c for c in lieferant if c.isalnum() or c in ' -_')

            endung = Path(src).suffix
            neuer_name = f"{datum}{betrag}_{lieferant}_{d.get('dateiname', '')}"
            # Sonderzeichen entfernen
            neuer_name = ''.join(c for c in neuer_name
                                  if c.isalnum() or c in ' .-_')[:80] + endung

            ziel = basis / ordner / neuer_name
            try:
                shutil.copy2(src, ziel)
                kopiert += 1
            except Exception:
                nicht_gefunden.append(d.get('dateiname', 'Unbekannt'))

        # Berichte generieren
        berichte_ordner = str(basis / '01_Berichte')
        self.export_pdf(
            dokumente,
            os.path.join(berichte_ordner, f'Steuerbericht_{steuerjahr}.pdf'),
            steuerjahr, statistiken
        )
        self.export_excel(
            dokumente,
            os.path.join(berichte_ordner, f'Steuerübersicht_{steuerjahr}.xlsx'),
            steuerjahr, statistiken
        )

        # Protokoll
        protokoll_pfad = str(basis / f'00_README_{steuerjahr}.txt')
        with open(protokoll_pfad, 'w', encoding='utf-8') as f:
            f.write(f"STEUERERKLÄRUNG {steuerjahr} – ÜBERGABEPAKET\n")
            f.write("=" * 60 + "\n\n")
            f.write(f"Erstellt am: {datetime.now().strftime('%d.%m.%Y %H:%M')}\n\n")
            f.write(f"Dokumente kopiert: {kopiert}\n")
            f.write(f"Gesamtbetrag: {statistiken.get('gesamtsumme', 0):.2f} EUR\n\n")
            f.write("ORDNERSTRUKTUR:\n")
            for ordner in sorted(unterordner.values()):
                anzahl = len(list((basis / ordner).iterdir()))
                f.write(f"  {ordner}/  ({anzahl} Dateien)\n")
            if nicht_gefunden:
                f.write("\nNICHT GEFUNDENE DATEIEN:\n")
                for nf in nicht_gefunden:
                    f.write(f"  - {nf}\n")
            f.write("\nHINWEIS: Bitte alle Dokumente auf Vollständigkeit prüfen.\n")
            f.write("Dieses Paket wurde von SteuerApp automatisch erstellt.\n")

        return str(basis)
