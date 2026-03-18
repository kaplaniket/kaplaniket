"""
Dokument-Prozessor: Extrahiert Text aus PDFs, Bildern und anderen Dateien
Nutzt PyMuPDF, Pillow und pytesseract (mit Fallback wenn nicht installiert)
"""

import os
import re
import io
import subprocess
from pathlib import Path
from typing import Optional, Tuple


# Optionale Imports mit Fallback
try:
    import fitz  # PyMuPDF
    HAS_PYMUPDF = True
except ImportError:
    HAS_PYMUPDF = False

try:
    from PIL import Image
    HAS_PIL = True
except ImportError:
    HAS_PIL = False

try:
    import pytesseract
    # Tesseract-Pfade für macOS
    for tp in ['/usr/local/bin/tesseract', '/opt/homebrew/bin/tesseract',
               '/usr/bin/tesseract']:
        if os.path.exists(tp):
            pytesseract.pytesseract.tesseract_cmd = tp
            break
    HAS_TESSERACT = True
except ImportError:
    HAS_TESSERACT = False

try:
    import docx
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


class DokumentProzessor:

    def __init__(self):
        self.ocr_sprachen = 'deu+eng'

    def extrahiere_text(self, dateipfad: str) -> Tuple[str, str]:
        """
        Extrahiert Text aus einer Datei.
        Gibt (text, methode) zurück.
        """
        pfad = Path(dateipfad)
        endung = pfad.suffix.lower()

        try:
            if endung == '.pdf':
                return self._verarbeite_pdf(dateipfad)
            elif endung in ['.jpg', '.jpeg', '.png', '.tiff', '.tif',
                            '.bmp', '.gif', '.heic', '.heif', '.webp']:
                return self._verarbeite_bild(dateipfad)
            elif endung in ['.docx']:
                return self._verarbeite_docx(dateipfad)
            elif endung in ['.xlsx', '.xls']:
                return self._verarbeite_excel(dateipfad)
            elif endung in ['.txt', '.csv', '.xml', '.html', '.htm', '.rtf']:
                return self._verarbeite_text(dateipfad)
            else:
                return '', 'nicht_unterstuetzt'
        except Exception as e:
            return '', f'fehler: {e}'

    def _verarbeite_pdf(self, pfad: str) -> Tuple[str, str]:
        text = ''
        methode = 'pdf_text'

        if HAS_PYMUPDF:
            try:
                doc = fitz.open(pfad)
                seiten_texte = []
                for seite in doc:
                    seiten_texte.append(seite.get_text())
                doc.close()
                text = '\n'.join(seiten_texte)
                methode = 'pymupdf'
            except Exception:
                pass

        # Fallback: OCR für gescannte PDFs (wenig/kein Text)
        if len(text.strip()) < 50 and HAS_PYMUPDF and HAS_TESSERACT and HAS_PIL:
            try:
                text = self._pdf_ocr(pfad)
                methode = 'pdf_ocr'
            except Exception:
                pass

        # Fallback: strings-Befehl (macOS/Linux)
        if not text.strip():
            try:
                result = subprocess.run(
                    ['strings', pfad],
                    capture_output=True, text=True, timeout=10
                )
                text = result.stdout
                methode = 'strings'
            except Exception:
                pass

        return text, methode

    def _pdf_ocr(self, pfad: str) -> str:
        """PDF zu Bild → OCR"""
        if not HAS_PYMUPDF or not HAS_TESSERACT or not HAS_PIL:
            return ''

        texte = []
        doc = fitz.open(pfad)
        for seite in doc:
            mat = fitz.Matrix(2, 2)  # 2x Zoom für bessere OCR-Qualität
            pix = seite.get_pixmap(matrix=mat)
            img_data = pix.tobytes("png")
            img = Image.open(io.BytesIO(img_data))
            t = pytesseract.image_to_string(img, lang=self.ocr_sprachen)
            texte.append(t)
        doc.close()
        return '\n'.join(texte)

    def _verarbeite_bild(self, pfad: str) -> Tuple[str, str]:
        if not HAS_TESSERACT or not HAS_PIL:
            return '', 'ocr_nicht_verfuegbar'

        try:
            img = Image.open(pfad)
            # Großes Bild für bessere OCR-Qualität skalieren
            if max(img.size) > 4000:
                factor = 4000 / max(img.size)
                new_size = (int(img.width * factor), int(img.height * factor))
                img = img.resize(new_size, Image.LANCZOS)
            text = pytesseract.image_to_string(img, lang=self.ocr_sprachen)
            return text, 'tesseract_ocr'
        except Exception as e:
            return '', f'ocr_fehler: {e}'

    def _verarbeite_docx(self, pfad: str) -> Tuple[str, str]:
        if not HAS_DOCX:
            return '', 'python_docx_fehlt'
        try:
            doc = docx.Document(pfad)
            text = '\n'.join(p.text for p in doc.paragraphs)
            return text, 'python_docx'
        except Exception as e:
            return '', f'docx_fehler: {e}'

    def _verarbeite_excel(self, pfad: str) -> Tuple[str, str]:
        if not HAS_OPENPYXL:
            return '', 'openpyxl_fehlt'
        try:
            wb = openpyxl.load_workbook(pfad, read_only=True, data_only=True)
            zeilen = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    zeile = ' | '.join(str(c) for c in row if c is not None)
                    if zeile.strip():
                        zeilen.append(zeile)
            return '\n'.join(zeilen), 'openpyxl'
        except Exception as e:
            return '', f'excel_fehler: {e}'

    def _verarbeite_text(self, pfad: str) -> Tuple[str, str]:
        for encoding in ['utf-8', 'latin-1', 'cp1252', 'iso-8859-1']:
            try:
                with open(pfad, 'r', encoding=encoding, errors='replace') as f:
                    return f.read(50000), f'text_{encoding}'
            except Exception:
                continue
        return '', 'text_fehler'

    def erstelle_thumbnail(self, pfad: str, ziel_pfad: str,
                           groesse: Tuple[int, int] = (200, 280)) -> bool:
        """Erstellt ein Vorschaubild"""
        endung = Path(pfad).suffix.lower()
        try:
            if endung == '.pdf' and HAS_PYMUPDF and HAS_PIL:
                doc = fitz.open(pfad)
                seite = doc[0]
                mat = fitz.Matrix(0.5, 0.5)
                pix = seite.get_pixmap(matrix=mat)
                img = Image.open(io.BytesIO(pix.tobytes("png")))
                img.thumbnail(groesse)
                os.makedirs(os.path.dirname(ziel_pfad), exist_ok=True)
                img.save(ziel_pfad, 'PNG')
                doc.close()
                return True
            elif endung in ['.jpg', '.jpeg', '.png', '.tiff', '.tif',
                            '.bmp', '.gif'] and HAS_PIL:
                img = Image.open(pfad)
                img.thumbnail(groesse)
                os.makedirs(os.path.dirname(ziel_pfad), exist_ok=True)
                img.save(ziel_pfad, 'PNG')
                return True
        except Exception:
            pass
        return False


def verfuegbare_methoden() -> dict:
    return {
        'PyMuPDF':    HAS_PYMUPDF,
        'Pillow':     HAS_PIL,
        'Tesseract':  HAS_TESSERACT,
        'python-docx': HAS_DOCX,
        'openpyxl':   HAS_OPENPYXL,
    }
